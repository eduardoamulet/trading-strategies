"""Batch backtest desde un Excel de configuraciones (modo «Cargar backtesting file»).

Cada FILA del Excel es UNA configuración (los parámetros del motor). `run_batch` corre el backtest de
CADA config sobre el producto (tickers × fechas) con el Tipo de operación dado, y emite filas
GRANULARES — una por (config × ticker × fecha) — con Inversión Total / Ganancia Total, listas para
mostrar en tabla y para volcar a un Excel descargable.

Lógica PURA (sin Streamlit) → testeable headless. Replica EXACTAMENTE el mapeo del panel de señales
(app.py ~L1457-1619): fills→entry_at_ask/exit_at_bid/nbbo, sentinelas umbral/stop, resolución, DTE,
filtro de confirmación, y la salida colectiva (ROI/stop) por día sobre el conjunto de tickers.
"""
from __future__ import annotations

import os
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from io import BytesIO
from pathlib import Path

import signals_backtest as sbt
from portfolio_exit import apply_collective_exit

# Granularidad (etiqueta del Excel) → resolución del Downloader (downloader.py RES).
_RES_MAP = {"1 min": "1min", "30 seg": "30s", "15 seg": "15s",
            "1min": "1min", "30s": "30s", "15s": "15s"}

# Columnas de salida (en orden) del resultado granular.
RESULT_COLS = ["ID", "Fecha", "Ticker", "Operación", "Inversión Total", "Ganancia Total",
               "ROI %", "status", "error"]


def _yes(v) -> bool:
    """¿La celda representa un 'Sí'? (tolera Sí/Si/Yes/True/1/✓/✅)."""
    return str(v).strip().lower() in ("sí", "si", "yes", "true", "1", "x", "✓", "✅", "verdadero")


def _num(v, default: float = 0.0) -> float:
    try:
        return float(str(v).replace(",", ".").strip())
    except Exception:
        return float(default)


def map_config(row: dict) -> dict:
    """Mapea una fila del Excel (columna→valor) a los kwargs de `sbt.run_one` + extras (resolución,
    salida colectiva, hora de entrada). PURA, no corre nada. Tolera columnas faltantes (usa defaults).
    """
    def g(*names):
        for n in names:
            if n in row and row[n] is not None and str(row[n]).strip() != "":
                return row[n]
        return None

    fills = str(g("Modelo de fills") or "").lower()
    f2 = "fase 2" in fills
    ef = f2 or "fase 1" in fills                      # entry_at_ask = exit_at_bid (Fase 1 o Fase 2)

    crit = str(g("Criterio de selección de contrato") or "").lower()
    criterio = "itm_first" if "itm" in crit else "spread"

    dte_s = str(g("Vencimiento DTE") or "0").strip()
    auto_dte = "auto" in dte_s.lower()
    dte = 1 if dte_s.startswith("1") else 0

    filtro = str(g("Filtro confirmación 1ª vela", "Filtro confirmación") or "No filtrar").lower()
    confirm = "no filtrar" not in filtro
    flip = ("flip" in filtro) or ("vuelta" in filtro)

    umbral = _num(g("Umbral ROI (%) del ticker"), 1000.0) if _yes(g("Cerrar si Umbral ROI ticker")) else 100000.0
    stop = -abs(_num(g("Stop loss (%) del ticker"), 100.0)) if _yes(g("Cerrar si Stop loss ticker")) else -100000.0

    coll_profit = (_num(g("Umbral ROI colectivo (%)"), 0.0) / 100.0) if _yes(g("Cerrar si Umbral ROI colectivo")) else None
    coll_stop = (-abs(_num(g("Stop loss (%) colectivo", "Stop loss (%) de colectivo"), 0.0)) / 100.0) \
        if _yes(g("Cerrar si Stop loss colectivo")) else None

    resolution = _RES_MAP.get(str(g("Granularidad temporal") or "1 min").strip(), "1min")

    entry_lbl = str(g("Horario de entrada") or "").strip()
    entry_hora = None if (not entry_lbl or "alerta" in entry_lbl.lower()) else entry_lbl

    run_kwargs = dict(
        inversion=_num(g("Inversión ($)"), 1000.0),
        umbral_pct=umbral, stop_pct=stop,
        entry_at_ask=ef, exit_at_bid=ef, nbbo_timeline=f2,
        auto_dte=auto_dte, selection_criterion=criterio,
        refuerzo_loss_pct=_num(g("Umbral pérdida refuerzo (%)"), 50.0) / 100.0,
        refuerzo_max=int(_num(g("No. de veces a reforzar"), 3)),
        call_pct=_num(g("Inversión CALL (%)"), 50.0),
        search_window_min=_num(g("Ventana búsqueda contrato (min)"), 4.0),
        dte=dte, exit_hora=str(g("Horario de salida") or "16:00").strip(),
        confirm_candle=confirm, confirm_min_body_pct=_num(g("Cuerpo mínimo anti-doji (%)"), 0.0),
        flip_on_wrong_direction=flip,
        apply_refuerzo=_yes(g("Aplicar refuerzo")),
        cut_weak_confirmation=_yes(g("Cerrar si confirmación débil")),
    )
    return {"run_kwargs": run_kwargs, "resolution": resolution,
            "coll_profit_frac": coll_profit, "coll_stop_frac": coll_stop,
            "entry_hora": entry_hora, "hora": str(g("Hora (HH:MM)", "Hora") or "09:30").strip(),
            "id": str(g("ID") or "").strip()}


def read_configs(file_or_path, sheet: str = "Backtesting", header_row: int = 2) -> list:
    """Lee las filas de configuración del Excel → lista de dicts {columna: valor}. Salta filas sin ID."""
    from openpyxl import load_workbook
    wb = load_workbook(file_or_path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[header_row]]
    configs = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {hdr[i]: ws.cell(r, i + 1).value for i in range(len(hdr)) if hdr[i]}
        if not row.get("ID"):
            continue
        configs.append(row)
    return configs


def auto_workers(n_tickers: int, n_dates: int) -> int:
    """Nº de hilos de procesamiento ≈ **1 por cada (ticker, fecha)**: clamp(n_tickers · n_dates, 8, 64).
    Mínimo 8 (con pocos pares igual se paralelizan las configs entre sí); máximo 64 (más hilos no ayudan:
    los limita el GIL y el nº de cores, y saturan el disco/NBBO)."""
    return max(8, min(64, int(n_tickers) * int(n_dates)))


def auto_processes() -> int:
    """Nº de procesos hijos del multiproceso real = nº de cores − 1 (deja uno para la UI/SO), mínimo 2."""
    return max(2, (os.cpu_count() or 4) - 1)


def run_batch(dl, configs: list, tickers: list, dates: list, tipo: str,
              progress_cb=None, max_workers: int = None) -> list:
    """Corre el batch. Para cada config × fecha, corre TODOS los tickers (en paralelo), aplica la salida
    colectiva de esa config sobre el día, y emite una fila GRANULAR por (config, ticker, fecha).

    `configs` admite filas crudas del Excel (se mapean con map_config) o dicts ya mapeados. La salida
    colectiva se aplica por (config, fecha) sobre el conjunto de tickers (igual que el panel). El
    Downloader es compartido → la resolución se setea por grupo de configs con la MISMA resolución
    (en el sweep todas son «1 min», así que es un solo grupo, sin race).

    Devuelve lista de dicts con RESULT_COLS.
    """
    mapped = [c if (isinstance(c, dict) and "run_kwargs" in c) else map_config(c) for c in configs]
    tickers = [str(t).upper().strip() for t in tickers]
    dates = [str(d).strip() for d in dates]
    if not max_workers:
        max_workers = auto_workers(len(tickers), len(dates))   # ~1 hilo por (ticker, fecha)

    by_res = defaultdict(list)
    for ci, m in enumerate(mapped):
        by_res[m["resolution"]].append(ci)

    results = {}                                   # (ci, fecha, ticker) → dict de run_one
    total = len(mapped) * len(dates) * max(1, len(tickers))
    done = [0]

    for res, cidxs in by_res.items():
        dl.resolution = res
        tasks = [(ci, d, tk) for ci in cidxs for d in dates for tk in tickers]
        with ThreadPoolExecutor(max_workers=max(1, min(max_workers, len(tasks) or 1))) as ex:
            futs = {}
            for (ci, d, tk) in tasks:
                m = mapped[ci]
                spec = {"ticker": tk, "fecha": d, "hora": m["entry_hora"] or m["hora"], "tipo": tipo}
                futs[ex.submit(sbt.run_one, dl, spec, **m["run_kwargs"])] = (ci, d, tk)
            for f in as_completed(futs):
                ci, d, tk = futs[f]
                try:
                    results[(ci, d, tk)] = f.result()
                except Exception as e:                                            # noqa: BLE001
                    results[(ci, d, tk)] = {"status": "error", "iteration": None, "error": str(e)}
                done[0] += 1
                if progress_cb:
                    progress_cb(done[0], total)

    out = []
    for ci, m in enumerate(mapped):
        for d in dates:
            res_list = [(tk, results.get((ci, d, tk),
                                         {"status": "error", "iteration": None, "error": "sin resultado"}))
                        for tk in tickers]
            if m["coll_profit_frac"] is not None or m["coll_stop_frac"] is not None:
                try:
                    apply_collective_exit(
                        [{"fecha": d, "ticker": tk, "iteration": r.get("iteration")} for tk, r in res_list],
                        m["coll_profit_frac"], m["coll_stop_frac"], stop_require_multi=True)
                except Exception:                                                 # noqa: BLE001
                    pass
            for tk, r in res_list:
                it = r.get("iteration")
                if r.get("status") == "ok" and it is not None:
                    inv, gain = float(it.invest_total), float(it.gain_total)
                    out.append({"ID": m["id"], "Fecha": d, "Ticker": tk, "Operación": tipo,
                                "Inversión Total": round(inv, 2), "Ganancia Total": round(gain, 2),
                                "ROI %": round(gain / inv * 100, 2) if inv else 0.0,
                                "status": "ok", "error": "",
                                "exit_reason": getattr(it, "exit_reason", "") or ""})
                else:
                    out.append({"ID": m["id"], "Fecha": d, "Ticker": tk, "Operación": tipo,
                                "Inversión Total": None, "Ganancia Total": None, "ROI %": None,
                                "status": r.get("status", "error"), "error": str(r.get("error") or ""),
                                "exit_reason": ""})
    return out


def batch_totals(rows: list) -> dict:
    """Totales AGREGADOS sobre TODAS las filas granulares (para el panel de métricas tipo «Totales del
    backtest»): Σ Inversión, Σ Ganancia, capital final, ganadores/perdedores, win rate y conteo de
    motivos de salida. Suma sobre todos los (config × ticker × fecha) OK."""
    ok = [r for r in rows if r.get("status") == "ok"]
    inv = sum(r.get("Inversión Total") or 0.0 for r in ok)
    gain = sum(r.get("Ganancia Total") or 0.0 for r in ok)
    wins = sum(1 for r in ok if (r.get("Ganancia Total") or 0.0) > 0)
    losses = sum(1 for r in ok if (r.get("Ganancia Total") or 0.0) < 0)
    reasons: dict = {}
    for r in ok:
        er = r.get("exit_reason") or ""
        reasons[er] = reasons.get(er, 0) + 1
    dec = wins + losses
    return {"n_ok": len(ok), "n_total": len(rows), "inv": round(inv, 2), "gain": round(gain, 2),
            "capital": round(inv + gain, 2), "roi": (gain / inv) if inv else 0.0,
            "wins": wins, "losses": losses, "win_rate": (wins / dec) if dec else 0.0,
            "reasons": reasons}


# ───────────────── MULTIPROCESO real (1 proceso hijo por core; tarea = config × fecha) ─────────────────
# Los hilos no aceleran (run_one es CPU-bound → GIL). Con procesos cada hijo corre en su propio
# intérprete → paralelismo real. Tarea por (config, fecha): el hijo corre TODOS los tickers de ese día
# y aplica la salida colectiva ahí dentro, así sólo viajan ESCALARES de vuelta (no DataFrames).
_MP_DL = None   # Downloader por-proceso (creado 1 vez en el initializer del Pool; no se puede picklear).


def _mp_init(data_dir: str, api_key: str) -> None:
    global _MP_DL
    from adapter_polygon import PolygonAdapter
    from downloader import Downloader
    _MP_DL = Downloader(PolygonAdapter(api_key, rate_limit_per_min=600), Path(data_dir))


def _mp_run_config_date(task: tuple) -> list:
    """task = (m, d, tickers, tipo). Corre los tickers de (config m, fecha d), aplica la colectiva y
    devuelve filas escalares (RESULT_COLS + exit_reason)."""
    m, d, tickers, tipo = task
    dl = _MP_DL
    dl.resolution = m["resolution"]
    res_list = []
    for tk in tickers:
        spec = {"ticker": tk, "fecha": d, "hora": m["entry_hora"] or m["hora"], "tipo": tipo}
        try:
            r = sbt.run_one(dl, spec, **m["run_kwargs"])
        except Exception as e:   # noqa: BLE001
            r = {"status": "error", "iteration": None, "error": str(e)}
        res_list.append((tk, r))
    if m["coll_profit_frac"] is not None or m["coll_stop_frac"] is not None:
        try:
            apply_collective_exit(
                [{"fecha": d, "ticker": tk, "iteration": r.get("iteration")} for tk, r in res_list],
                m["coll_profit_frac"], m["coll_stop_frac"], stop_require_multi=True)
        except Exception:   # noqa: BLE001
            pass
    out = []
    for tk, r in res_list:
        it = r.get("iteration")
        if r.get("status") == "ok" and it is not None:
            inv, gain = float(it.invest_total), float(it.gain_total)
            out.append({"ID": m["id"], "Fecha": d, "Ticker": tk, "Operación": tipo,
                        "Inversión Total": round(inv, 2), "Ganancia Total": round(gain, 2),
                        "ROI %": round(gain / inv * 100, 2) if inv else 0.0,
                        "status": "ok", "error": "", "exit_reason": getattr(it, "exit_reason", "") or ""})
        else:
            out.append({"ID": m["id"], "Fecha": d, "Ticker": tk, "Operación": tipo,
                        "Inversión Total": None, "Ganancia Total": None, "ROI %": None,
                        "status": r.get("status", "error"), "error": str(r.get("error") or ""), "exit_reason": ""})
    return out


def run_batch_parallel(data_dir, api_key, configs: list, tickers: list, dates: list, tipo: str,
                       progress_cb=None, processes: int = None) -> list:
    """Igual que run_batch pero con MULTIPROCESO real (sin límite del GIL). Un proceso hijo por core;
    cada tarea = (config × fecha) corre todos los tickers de ese día. Devuelve lista de dicts."""
    import multiprocessing as mp
    mapped = [c if (isinstance(c, dict) and "run_kwargs" in c) else map_config(c) for c in configs]
    tickers = [str(t).upper().strip() for t in tickers]
    dates = [str(d).strip() for d in dates]
    if not processes:
        processes = auto_processes()
    tasks = [(m, d, tickers, tipo) for m in mapped for d in dates]
    processes = max(1, min(processes, len(tasks)))
    total_bt = len(tasks) * max(1, len(tickers))
    out, done = [], 0
    with mp.Pool(processes=processes, initializer=_mp_init, initargs=(str(data_dir), api_key)) as pool:
        for rows in pool.imap_unordered(_mp_run_config_date, tasks):
            out.extend(rows)
            done += len(tickers)
            if progress_cb:
                progress_cb(min(done, total_bt), total_bt)
    return out


def summarize(rows: list) -> list:
    """Agrega las filas granulares POR CONFIG (ID): suma Inversión/Ganancia de los backtests OK,
    ROI% combinado, y conteos ok/error. Mantiene el orden de aparición. Vista para comparar configs."""
    agg, order = {}, []
    for r in rows:
        i = r.get("ID", "")
        if i not in agg:
            agg[i] = {"ID": i, "Inversión Total": 0.0, "Ganancia Total": 0.0, "ok": 0, "error": 0, "n": 0}
            order.append(i)
        a = agg[i]
        a["n"] += 1
        if r.get("status") == "ok":
            a["ok"] += 1
            a["Inversión Total"] += r.get("Inversión Total") or 0.0
            a["Ganancia Total"] += r.get("Ganancia Total") or 0.0
        else:
            a["error"] += 1
    out = []
    for i in order:
        a = agg[i]
        inv = a["Inversión Total"]
        a["ROI %"] = round(a["Ganancia Total"] / inv * 100, 2) if inv else 0.0
        a["Inversión Total"] = round(inv, 2)
        a["Ganancia Total"] = round(a["Ganancia Total"], 2)
        out.append(a)
    return out


def results_to_xlsx(rows: list, meta: dict | None = None) -> bytes:
    """Construye el Excel de resultados (bytes para st.download_button). 2 hojas:
      • «Resumen por config» — un renglón por config (ID) con Inversión/Ganancia/ROI sumados.
      • «Detalle» — un renglón por (config × ticker × fecha) con Fecha/Ticker/Operación + Inversión/Ganancia.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen por config"
    scols = ["ID", "Inversión Total", "Ganancia Total", "ROI %", "ok", "error", "n"]
    for ci, c in enumerate(scols, 1):
        cell = ws.cell(1, ci, c)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E2EFDA")
        cell.alignment = Alignment(horizontal="center")
    for ri, row in enumerate(summarize(rows), 2):
        for ci, c in enumerate(scols, 1):
            ws.cell(ri, ci, row.get(c))
    ws.freeze_panes = "A2"
    for ci in range(1, len(scols) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 16

    wd = wb.create_sheet("Detalle")
    for ci, c in enumerate(RESULT_COLS, 1):
        cell = wd.cell(1, ci, c)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(horizontal="center")
    for ri, row in enumerate(rows, 2):
        for ci, c in enumerate(RESULT_COLS, 1):
            wd.cell(ri, ci, row.get(c))
    wd.freeze_panes = "A2"
    for ci, c in enumerate(RESULT_COLS, 1):
        wd.column_dimensions[get_column_letter(ci)].width = 50 if c == "error" else 15

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
