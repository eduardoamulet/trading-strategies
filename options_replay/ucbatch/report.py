"""Excel Report Generator — workbook de salida con el formato del template.

Independiente del engine: recibe `seed` + `rows` (dicts escalares del runner) y arma el .xlsx.
Layout (hoja «Backtesting Results») — header de UNA fila → columnas PLANAS y ORDENABLES (Excel
Data→Sort/Filter necesita el header en la fila 1, sin celdas combinadas):
  r1: ID | Ticker | Fecha | Inversión total | Ganancia | Valor mínimo ROI (%/$) | Valor máximo ROI
      (%/$) | Promedios de todos los ROI (%/$) | Número ROI(%) >0 | <=0 | Número de errores  (+ auto-filtro)
  r2+: una fila por (ticker × día × escenario), ordenadas por Ticker → Fecha → ID.
Antes el header ocupaba 3 filas con merges (no ordenable); se aplanó a 1 fila.
Naming: Backtesting_Results_<TIPO>_of_<T1_T2_…>_from_<inicio>_to_<fin>.xlsx (fechas con guión largo).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HFILL = PatternFill("solid", fgColor="DDEBF7")

# (clave-en-row, decimales) en orden de columna de datos (col D=4 en adelante).
_DATA = [("inversion", 0), ("ganancia", 2),
         ("roi_min_pct", 2), ("roi_min_usd", 2), ("roi_max_pct", 2), ("roi_max_usd", 2),
         ("roi_avg_pct", 2), ("roi_avg_usd", 2), ("n_roi_pos", 0), ("n_roi_neg", 0), ("n_err", 0)]

# Header PLANO de una sola fila (fila 1) → columnas ordenables. El ORDEN mapea 1:1 con [ID, Ticker,
# Fecha] + _DATA (col D..N), así que fill_results/build escriben por posición sin ambigüedad.
_HEADERS = ["ID", "Ticker", "Fecha", "Inversión total", "Ganancia",
            "Valor mínimo ROI (%)", "Valor mínimo ROI ($)",
            "Valor máximo ROI (%)", "Valor máximo ROI ($)",
            "Promedios de todos los ROI (%)", "Promedios de todos los ROI ($)",
            "Número ROI(%) > 0", "Número ROI(%) <= 0", "Número de errores"]

# Columnas ENRIQUECIDAS (solo en el output DETALLADO `build`/`write`, 1 fila = 1 posición): la señal
# del Market Direction Engine a la entrada + el snapshot del contrato. (etiqueta, clave-en-row).
_ENRICH = [
    ("Dir · Acción", "md_action"), ("Dir · Score", "md_score"),
    ("Dir · Confianza", "md_confidence"), ("Dir · Trend", "md_trend"),
    ("Modo", "mode"), ("Strike CALL", "call_strike"), ("Strike PUT", "put_strike"),
    ("CALL bid", "call_bid"), ("CALL ask", "call_ask"), ("CALL spread", "call_spread"),
    ("PUT bid", "put_bid"), ("PUT ask", "put_ask"), ("PUT spread", "put_spread"),
    ("Prima CALL", "call_entry_prem"), ("Prima PUT", "put_entry_prem"),
    ("Spot entrada", "spot_at_start"), ("Motivo salida", "exit_reason"),
    ("Duración (min)", "duration_min"), ("OCC CALL", "call_occ"), ("OCC PUT", "put_occ"),
    ("CALL Delta", "call_delta"), ("CALL Gamma", "call_gamma"), ("CALL Theta", "call_theta"),
    ("CALL IV", "call_iv"), ("PUT Delta", "put_delta"), ("PUT Gamma", "put_gamma"),
    ("PUT Theta", "put_theta"), ("PUT IV", "put_iv"),
]


def output_filename(seed, tag: str = "") -> str:
    """`tag` (opcional) = slug de la COMBINACIÓN — evita que dos combinaciones con el mismo
    rango/tickers se pisen el results file."""
    tipo = seed.tipo.replace(" ", "_")
    tks = "_".join(seed.tickers)
    fi = seed.fecha_inicial.replace("-", "–")          # restituir guión largo (como el template)
    ff = seed.fecha_final.replace("-", "–")
    _tag = f"_{str(tag).strip().replace(' ', '_')}" if str(tag or "").strip() else ""
    return f"Backtesting_Results{_tag}_{tipo}_of_{tks}_from_{fi}_to_{ff}.xlsx"


def _headers(ws, enrich: bool = False) -> None:
    """Escribe el header PLANO en la fila 1 → columnas ordenables (Excel Data→Sort/Filter). Con
    `enrich=True` agrega las columnas enriquecidas (señal de dirección + snapshot del contrato)."""
    names = list(_HEADERS) + ([lbl for lbl, _ in _ENRICH] if enrich else [])
    for i, name in enumerate(names, 1):
        cell = ws.cell(1, i)
        cell.value = name
        cell.font = Font(bold=True, size=9)
        cell.alignment = _CTR
        cell.fill = _HFILL
    ws.row_dimensions[1].height = 30


def _copy_listas(wb, listas_src) -> None:
    """Copia la hoja «Listas» del template de entrada (dropdowns), si se provee."""
    if not listas_src:
        return
    try:
        from openpyxl import load_workbook
        src = load_workbook(Path(listas_src), read_only=True, data_only=True)
        if "Listas" in src.sheetnames:
            ws = wb.create_sheet("Listas")
            for r in src["Listas"].iter_rows(values_only=True):
                ws.append(list(r))
    except Exception:
        pass


def build(seed, rows: list, listas_src=None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Backtesting Results"
    # Enriquecido si las filas traen los campos extra (señal de dirección / snapshot del contrato).
    _enr = any(k in r for r in rows[:1] for _, k in _ENRICH) if rows else False
    _headers(ws, enrich=_enr)                           # header plano en la fila 1
    _tk = {t: i for i, t in enumerate(seed.tickers)}   # orden del seed (QQQ→SPY→IWM), no alfabético
    ordered = sorted(rows, key=lambda r: (_tk.get(str(r.get("Ticker")), 999), str(r.get("Fecha")), str(r.get("ID"))))
    for r in ordered:
        vals = [r.get("ID"), r.get("Ticker"), r.get("Fecha")]
        for key, dec in _DATA:
            v = r.get(key)
            vals.append(round(v, dec) if isinstance(v, (int, float)) else v)
        if _enr:
            for _, key in _ENRICH:
                v = r.get(key)
                vals.append(round(v, 4) if isinstance(v, float) else v)
        ws.append(vals)
    ws.freeze_panes = "A2"                              # fija el header (1 fila)
    _ncols = len(_HEADERS) + (len(_ENRICH) if _enr else 0)
    ws.auto_filter.ref = f"A1:{get_column_letter(_ncols)}{max(1, ws.max_row)}"   # orden/filtro 1-clic
    widths = [8, 8, 14, 13, 11, 17, 17, 17, 17, 20, 20, 15, 15, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _copy_listas(wb, listas_src)
    return wb


def write(seed, rows: list, out_dir, listas_src=None, tag: str = "") -> Path:
    """Arma y guarda el workbook. Devuelve el Path."""
    wb = build(seed, rows, listas_src=listas_src)
    out = Path(out_dir) / output_filename(seed, tag=tag)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ── Modo «rellenar el results file provisto»: 1 fila por ESCENARIO (agrega sus runs) ──────────
def aggregate_by_scenario(rows: list) -> dict:
    """Agrega las filas por-posición → resumen por ID usando el ROI FINAL de cada run.
    Valor mín/máx = peor/mejor run · Promedios = media de los runs · Inversión/Ganancia = sumas."""
    from collections import defaultdict
    groups: dict = defaultdict(list)
    for r in rows:
        groups[str(r.get("ID"))].append(r)
    out = {}
    for sid, rs in groups.items():
        inv_tot = sum(float(r.get("inversion") or 0.0) for r in rs)
        gan_tot = sum(float(r.get("ganancia") or 0.0) for r in rs)
        ok = [r for r in rs if not r.get("n_err")]
        usds = [float(r.get("ganancia") or 0.0) for r in ok]
        pcts = [float(r["ganancia"]) / float(r["inversion"]) * 100.0
                for r in ok if float(r.get("inversion") or 0.0) > 0]
        out[sid] = {
            "inversion": inv_tot, "ganancia": gan_tot,
            "roi_min_pct": min(pcts) if pcts else 0.0, "roi_min_usd": min(usds) if usds else 0.0,
            "roi_max_pct": max(pcts) if pcts else 0.0, "roi_max_usd": max(usds) if usds else 0.0,
            "roi_avg_pct": (sum(pcts) / len(pcts)) if pcts else 0.0,
            "roi_avg_usd": (sum(usds) / len(usds)) if usds else 0.0,
            "n_roi_pos": sum(1 for u in usds if u > 0),
            "n_roi_neg": sum(1 for u in usds if u <= 0),
            "n_err": sum(int(r.get("n_err") or 0) for r in rs),
        }
    return out


def fill_results(seed, rows: list, template_path, out_path) -> Path:
    """Rellena el results file provisto (hoja «Backtesting Results»): 1 fila por escenario,
    matcheada por ID en la col A. ROBUSTO al layout del header: rellena CUALQUIER fila cuyo col A sea
    un ID de escenario conocido → funciona igual con el header nuevo de 1 fila (datos en r2+) que con
    el viejo de 3 (datos en r4+), sin hardcodear la fila de inicio. Preserva formato. Devuelve el Path."""
    from openpyxl import load_workbook
    agg = aggregate_by_scenario(rows)
    tickers_lbl = ", ".join(seed.tickers)
    fecha_lbl = f"{seed.fecha_inicial} → {seed.fecha_final}"
    wb = load_workbook(Path(template_path))
    ws = wb["Backtesting Results"] if "Backtesting Results" in wb.sheetnames else wb.active
    _decs = [0, 2, 2, 2, 2, 2, 2, 2, 0, 0, 0]   # D..N (mismo orden que _DATA)
    keys = [k for k, _ in _DATA]
    n_filled = 0
    for row in ws.iter_rows(min_row=1):
        sid = str(row[0].value or "").strip()
        a = agg.get(sid)                    # None si col A no es un ID conocido (headers/filas vacías)
        if a is None:
            continue
        row[1].value = tickers_lbl          # B · Ticker (alcance)
        row[2].value = fecha_lbl            # C · Fecha (rango)
        for i, (key, dec) in enumerate(zip(keys, _decs)):
            v = a.get(key)
            row[3 + i].value = round(v, dec) if isinstance(v, (int, float)) else v
        n_filled += 1
    if not n_filled:                        # ningún ID matcheó → header/columna A distinta: avisar, no fallar mudo
        import obs_log
        obs_log.get_logger("ucbatch").warning(
            "fill_results: 0 filas rellenadas — ¿la col A del results file no tiene los IDs (C001…)? "
            "IDs esperados: %s", sorted(agg)[:5])
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
