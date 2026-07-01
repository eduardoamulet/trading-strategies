"""Excel Reader — lee `Backtesting_use_cases_template.xlsx`: Data seed (globales) + scenarios.

Aislado: solo I/O de Excel → estructuras de datos planas (`Seed`, `Scenario`). Sin lógica de negocio,
sin dependencia de la UI ni del engine. Robusto a:
  • header duplicado «Fecha inicial» (toma las dos columnas «Fecha…» por orden: inicial, final),
  • guión largo en las fechas (2026–02–01 → 2026-02-01),
  • columnas reordenadas (mapea por substring del header, no por posición fija).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

SEED_SHEET = "Data seed"
SCEN_SHEET = "Backtesting scenarios"
HEADER_ROW = 2          # fila 1 = título de sección; fila 2 = headers; fila 3+ = datos


def _norm_date(v) -> str:
    """«2026–02–01» / «2026—02—01» / datetime → «2026-02-01»."""
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).replace("–", "-").replace("—", "-").strip()


def _find(headers: list, *subs: str):
    """Índice de la 1ª columna cuyo header contiene TODOS los `subs` (case-insensitive)."""
    for i, h in enumerate(headers):
        hl = str(h or "").lower()
        if all(s.lower() in hl for s in subs):
            return i
    return None


def _find_all(headers: list, sub: str) -> list:
    return [i for i, h in enumerate(headers) if sub.lower() in str(h or "").lower()]


def _get(vals: list, idx):
    return vals[idx] if (idx is not None and idx < len(vals)) else None


@dataclass
class Seed:
    """Parámetros GLOBALES (compartidos por todos los escenarios)."""
    tickers: list           # ["QQQ", "SPY", "IWM"]
    tipo: str               # "CALL y PUT"
    fecha_inicial: str      # "2026-02-01"
    fecha_final: str        # "2026-02-03"
    inversion: float
    call_pct: float
    put_pct: float
    entrada: str            # "09:30"
    salida: str             # "13:55"
    ventana_min: float
    criterio: str           # texto crudo ("Menor spread en rango óptimo")
    fills: str              # texto crudo ("NBBO por barra · triggers sobre el bid (Fase 2)")
    granularidad_seg: int   # 60 (= cada cuánto se chequea el ROI / resolución de barras)
    dte: str                # "0 — mismo día"
    display: list = field(default_factory=list)   # [(label, value)] para mostrar en la UI


@dataclass
class Scenario:
    """Un escenario: ID + condiciones de entrada/salida (dict crudo header→valor)."""
    id: str
    cond: dict = field(default_factory=dict)


def read_seed(ws) -> Seed:
    rows = list(ws.iter_rows(min_row=HEADER_ROW, values_only=True))
    headers = list(rows[0]) if rows else []
    vals = list(rows[1]) if len(rows) > 1 else []
    fechas = _find_all(headers, "fecha")
    fi = fechas[0] if fechas else None
    ff = fechas[1] if len(fechas) > 1 else None
    tickers_raw = str(_get(vals, _find(headers, "ticker")) or "")
    seed = Seed(
        tickers=[t.strip().upper() for t in tickers_raw.split(",") if t.strip()],
        tipo=str(_get(vals, _find(headers, "tipo")) or "").strip(),
        fecha_inicial=_norm_date(_get(vals, fi)),
        fecha_final=_norm_date(_get(vals, ff)),
        inversion=float(_get(vals, _find(headers, "inversión", "($)")) or _get(vals, _find(headers, "inversion ($")) or 1000),
        call_pct=float(_get(vals, _find(headers, "call")) or 50),
        put_pct=float(_get(vals, _find(headers, "put")) or 50),
        entrada=str(_get(vals, _find(headers, "entrada")) or "09:30").strip(),
        salida=str(_get(vals, _find(headers, "salida")) or "16:00").strip(),
        ventana_min=float(_get(vals, _find(headers, "ventana")) or 0),
        criterio=str(_get(vals, _find(headers, "criterio")) or "").strip(),
        fills=str(_get(vals, _find(headers, "fills") if _find(headers, "fills") is not None else _find(headers, "modelo")) or "").strip(),
        granularidad_seg=int(float(_get(vals, _find(headers, "granularidad") if _find(headers, "granularidad") is not None else _find(headers, "segundos")) or 60)),
        dte=str(_get(vals, _find(headers, "dte") if _find(headers, "dte") is not None else _find(headers, "vencimiento")) or "0 — mismo día").strip(),
    )
    # Para la UI: pares (header, valor) tal cual, saltando headers vacíos.
    seed.display = [(str(h), ("" if v is None else v)) for h, v in zip(headers, vals) if h]
    return seed


def read_scenarios(ws) -> list:
    rows = list(ws.iter_rows(min_row=HEADER_ROW, values_only=True))
    if not rows:
        return []
    headers = [str(h or "").strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or not r[0] or not str(r[0]).strip():
            continue
        cond = {h: v for h, v in zip(headers, r) if h}
        out.append(Scenario(id=str(r[0]).strip(), cond=cond))
    return out


def read_template(path) -> tuple:
    """Devuelve (Seed, list[Scenario]). Lanza si faltan las hojas requeridas."""
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    if SEED_SHEET not in wb.sheetnames or SCEN_SHEET not in wb.sheetnames:
        raise ValueError(f"El Excel debe tener las hojas «{SEED_SHEET}» y «{SCEN_SHEET}». "
                         f"Encontradas: {wb.sheetnames}")
    return read_seed(wb[SEED_SHEET]), read_scenarios(wb[SCEN_SHEET])
