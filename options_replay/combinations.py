"""Combinaciones de Backtesting — reemplaza el template estático de 480 escenarios.

Cada COMBINACIÓN nace de un Excel de variables (hoja «Backtesting variables»: columnas =
variables, filas = valores posibles). Sus ESCENARIOS son el producto cartesiano completo de
los valores de las columnas de CONDICIONES; las columnas de SEED (tickers, fechas, inversión,
horarios, fills…) son globales de la corrida — igual que el «Data seed» del template viejo.

Persistencia: data/combinations.db (SQLite):
  · combinations : id, nombre, archivo, seed_json, variables_json, n_escenarios, estado…
  · scenarios    : (combination_id, codigo) → cond_json con los HEADERS ORIGINALES del
                   template — así `ucbatch.scenario.map_scenario` y el motor de interpretación
                   (bt_analysis) los consumen SIN cambios.
  · config       : combinación ACTIVA (la que usan el job diario y la reevaluación).

Adaptadores (la costura con el pipeline existente):
  · seed_for()           → ucbatch.reader.Seed (con override de fechas/tickers por corrida)
  · scenarios_for_batch()→ list[ucbatch.reader.Scenario]
  · scenarios_df()       → DataFrame canónico (drop-in de bt_analysis.loader.load_template()[1])
  · ensure_legacy()      → importa el template 480 UNA vez como combinación «tpl480» (después
                           de eso el pipeline no vuelve a depender del archivo).

Escala: la generación es un stream (itertools.product + executemany por lotes) — no materializa
la lista completa en RAM; `expected_scenarios` avisa el tamaño ANTES de generar.
"""
from __future__ import annotations

import itertools
import json
import sqlite3
from contextlib import closing
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).parent
DB_PATH = HERE / "data" / "combinations.db"
LEGACY_ID = "tpl480"
LEGACY_TEMPLATE = HERE.parent / "excels for backtesting" / "Backtesting_use_cases_template.xlsx"

VARS_SHEET = "Backtesting variables"

# Columnas de SEED (globales de la corrida) — mismas del «Data seed» del template viejo.
SEED_COLS = ["Tickers", "Tipo de operacion", "Fecha inicial", "Fecha inicial",   # 2ª = final
             "Inversión ($)", "Inversión CALL (%)", "Inversión PUT (%)",
             "Horario de entrada", "Horario de salida", "Ventana búsqueda contrato (max)",
             "Criterio de selección de contrato", "Modelo de fills",
             "Granularidad temporal (segundos)", "Vencimiento DTE"]

# Columnas de CONDICIONES (variables del producto cartesiano) — headers EXACTOS del template
# (los mismos que consumen map_scenario y loader._SCENARIO_RENAME).
COND_COLS = ["Alcance de salida", "Aplicar refuerzo", "Umbral pérdida refuerzo (%)",
             "No. de veces a reforzar", "Cerrar si Umbral ROI ticker", "Umbral ROI (%) del ticker",
             "Cerrar si Stop loss ticker", "Stop loss (%) del ticker",
             "Filtro confirmación 1ª vela", "Cerrar si confirmación débil",
             "Cuerpo mínimo anti-doji (%)", "Cerrar si Umbral ROI colectivo",
             "Umbral ROI colectivo (%)", "Cerrar si Stop loss colectivo",
             "Stop loss (%) colectivo"]

# Columnas de condiciones OPCIONALES: entran al producto cartesiano SOLO si el Excel las trae
# (los archivos viejos sin ellas siguen siendo válidos — no rompen la validación).
# «Tipo de operación (escenario)»: permite comparar POLÍTICAS DE SALIDA (CALL y PUT / CALL o
# PUT / plus / End of Day / Until reach ROI(%)) como una variable más; vacía o ausente → se
# usa el «Tipo de operacion» global del seed, como siempre.
COND_COLS_OPT = ["Tipo de operación (escenario)"]


# ── Infraestructura ───────────────────────────────────────────────────────────
def _connect(path: Path = DB_PATH) -> sqlite3.Connection:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(path))
    con.execute("""CREATE TABLE IF NOT EXISTS combinations (
        id TEXT PRIMARY KEY, nombre TEXT NOT NULL, archivo TEXT, creado_en TEXT NOT NULL,
        estado TEXT NOT NULL DEFAULT 'importada', seed_json TEXT NOT NULL,
        variables_json TEXT NOT NULL, n_escenarios INTEGER NOT NULL DEFAULT 0,
        generado_en TEXT)""")
    con.execute("""CREATE TABLE IF NOT EXISTS scenarios (
        combination_id TEXT NOT NULL, codigo TEXT NOT NULL, cond_json TEXT NOT NULL,
        creado_en TEXT NOT NULL, PRIMARY KEY (combination_id, codigo))""")
    con.execute("CREATE TABLE IF NOT EXISTS config (clave TEXT PRIMARY KEY, valor TEXT)")
    return con


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _norm(v) -> str:
    """Celda → string canónico: fechas con guión largo normalizadas, resto strip."""
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).replace("–", "-").replace("—", "-").strip()


# ── Importación ───────────────────────────────────────────────────────────────
def import_file(src, *, nombre: str | None = None, path: Path = DB_PATH) -> dict:
    """Importa un Excel de variables como una COMBINACIÓN nueva (nunca sobrescribe: el id lleva
    timestamp). Valida la estructura (hoja + columnas de condiciones completas + seed mínimo),
    extrae los valores únicos por columna (en orden de aparición) y persiste todo.
    `src` = path o file-like (UploadedFile). Devuelve el registro creado."""
    from openpyxl import load_workbook

    # SIN read_only: con estos files openpyxl read-only devuelve None en celdas con valor
    # (quirk observado con los xlsx de variables); son chicos (~500 filas) — modo normal.
    if hasattr(src, "read"):
        try:
            src.seek(0)
        except Exception:  # noqa: BLE001
            pass
        wb = load_workbook(src, data_only=True)
        archivo = getattr(src, "name", "combinacion.xlsx")
    else:
        wb = load_workbook(Path(src), data_only=True)
        archivo = Path(src).name
    if VARS_SHEET not in wb.sheetnames:
        raise ValueError(f"El Excel debe tener la hoja «{VARS_SHEET}». "
                         f"Encontradas: {wb.sheetnames}")
    rows = list(wb[VARS_SHEET].iter_rows(values_only=True))
    if not rows:
        raise ValueError("La hoja de variables está vacía.")
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]

    # Validación de estructura: TODAS las condiciones presentes + seed mínimo.
    faltan = [c for c in COND_COLS if c not in hdr]
    if faltan:
        raise ValueError("Estructura inválida — faltan columnas de condiciones: "
                         + ", ".join(faltan))
    for req in ("Tickers", "Horario de entrada", "Horario de salida"):
        if req not in hdr:
            raise ValueError(f"Estructura inválida — falta la columna de seed «{req}».")

    def _col_values(j: int) -> list[str]:
        vals, seen = [], set()
        for r in rows[1:]:
            v = _norm(r[j] if j < len(r) else None)
            if v and v not in seen:
                seen.add(v)
                vals.append(v)
        return vals

    # SEED: primera aparición de cada columna de seed (las dos «Fecha inicial» por orden);
    # Tickers = TODOS los valores de la columna (lista, como el seed del template).
    fecha_idxs = [j for j, h in enumerate(hdr) if h == "Fecha inicial"]
    seed: dict = {}
    for j, h in enumerate(hdr):
        if h in ("", "Fecha inicial") or h in COND_COLS or h in COND_COLS_OPT:
            continue
        vals = _col_values(j)
        if h == "Tickers":
            seed["tickers"] = [t.upper() for t in vals]
        elif h not in seed:
            seed[h] = vals[0] if vals else ""
    seed["fecha_inicial"] = _col_values(fecha_idxs[0])[0] if fecha_idxs else ""
    seed["fecha_final"] = (_col_values(fecha_idxs[1])[0]
                           if len(fecha_idxs) > 1 else seed.get("fecha_inicial", ""))
    if not seed.get("tickers"):
        raise ValueError("La columna «Tickers» no tiene valores.")

    # VARIABLES del producto cartesiano: valores únicos por columna de condición (orden hoja).
    # Una columna VACÍA es legítima (p. ej. «solo tickers» deja las de colectivo en blanco):
    # cuenta como UN valor vacío → multiplicador 1 en el producto, y el motor le aplica sus
    # defaults (flag vacío = No; número vacío = default del batch) — igual que el template.
    variables: dict = {}
    for c in COND_COLS:
        variables[c] = _col_values(hdr.index(c)) or [""]
    for c in COND_COLS_OPT:                      # opcionales: solo si el Excel trae la columna
        if c in hdr:
            variables[c] = _col_values(hdr.index(c)) or [""]

    cid = "comb_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:22]
    reg = {"id": cid, "nombre": (nombre or archivo).strip(), "archivo": archivo,
           "creado_en": _now(), "estado": "importada", "seed": seed, "variables": variables,
           "n_escenarios": 0, "generado_en": None}
    with closing(_connect(path)) as con, con:
        con.execute("INSERT INTO combinations (id, nombre, archivo, creado_en, estado, "
                    "seed_json, variables_json, n_escenarios) VALUES (?,?,?,?,?,?,?,0)",
                    (cid, reg["nombre"], archivo, reg["creado_en"], "importada",
                     json.dumps(seed, ensure_ascii=False),
                     json.dumps(variables, ensure_ascii=False)))
    return reg


# ── Generación de escenarios (producto cartesiano, streaming) ─────────────────
def expected_scenarios(combo: dict | str, path: Path = DB_PATH) -> int:
    """Tamaño del producto cartesiano SIN generarlo (para avisar antes)."""
    c = combo if isinstance(combo, dict) else get_combination(combo, path=path)
    n = 1
    for vals in (c.get("variables") or {}).values():
        n *= max(1, len(vals))
    return n


def generate_scenarios(combo_id: str, *, batch: int = 5000, path: Path = DB_PATH,
                       progress_cb=None) -> int:
    """Genera TODOS los escenarios (producto cartesiano completo de las variables) y los
    persiste. Streaming: nunca materializa la lista completa en RAM; INSERT por lotes.
    Re-generar es idempotente (borra los anteriores de ESTA combinación). Devuelve el total."""
    c = get_combination(combo_id, path=path)
    if not c:
        raise ValueError(f"Combinación {combo_id!r} inexistente.")
    variables = c["variables"]
    cols = list(variables.keys())                      # orden de la hoja (COND_COLS)
    total = expected_scenarios(c)
    width = max(3, len(str(total)))
    now = _now()
    with closing(_connect(path)) as con, con:
        con.execute("DELETE FROM scenarios WHERE combination_id=?", (combo_id,))
        lote, n = [], 0
        for combo_vals in itertools.product(*(variables[c_] for c_ in cols)):
            n += 1
            cond = dict(zip(cols, combo_vals))
            lote.append((combo_id, f"C{n:0{width}d}", json.dumps(cond, ensure_ascii=False), now))
            if len(lote) >= batch:
                con.executemany("INSERT INTO scenarios VALUES (?,?,?,?)", lote)
                lote.clear()
                if progress_cb:
                    progress_cb(n, total)
        if lote:
            con.executemany("INSERT INTO scenarios VALUES (?,?,?,?)", lote)
        con.execute("UPDATE combinations SET n_escenarios=?, estado='generada', generado_en=? "
                    "WHERE id=?", (n, now, combo_id))
    return n


# ── Consultas / administración ────────────────────────────────────────────────
def _row_to_dict(r) -> dict:
    return {"id": r[0], "nombre": r[1], "archivo": r[2], "creado_en": r[3], "estado": r[4],
            "seed": json.loads(r[5]), "variables": json.loads(r[6]),
            "n_escenarios": int(r[7] or 0), "generado_en": r[8]}


def list_combinations(path: Path = DB_PATH) -> list[dict]:
    with closing(_connect(path)) as con:
        rows = con.execute("SELECT id, nombre, archivo, creado_en, estado, seed_json, "
                           "variables_json, n_escenarios, generado_en FROM combinations "
                           "ORDER BY creado_en").fetchall()
    return [_row_to_dict(r) for r in rows]


def get_combination(combo_id: str, path: Path = DB_PATH) -> dict | None:
    with closing(_connect(path)) as con:
        r = con.execute("SELECT id, nombre, archivo, creado_en, estado, seed_json, "
                        "variables_json, n_escenarios, generado_en FROM combinations "
                        "WHERE id=?", (combo_id,)).fetchone()
    return _row_to_dict(r) if r else None


def delete_combination(combo_id: str, path: Path = DB_PATH) -> None:
    """Borra la combinación y sus escenarios. La activa no se puede borrar (proteger el job)."""
    if active_combination(path=path) == combo_id:
        raise ValueError("La combinación ACTIVA no se puede eliminar — activá otra primero.")
    with closing(_connect(path)) as con, con:
        con.execute("DELETE FROM scenarios WHERE combination_id=?", (combo_id,))
        con.execute("DELETE FROM combinations WHERE id=?", (combo_id,))


def active_combination(path: Path = DB_PATH) -> str | None:
    with closing(_connect(path)) as con:
        r = con.execute("SELECT valor FROM config WHERE clave='active_combination'").fetchone()
    return r[0] if r else None


def set_active(combo_id: str, path: Path = DB_PATH) -> None:
    c = get_combination(combo_id, path=path)
    if not c:
        raise ValueError(f"Combinación {combo_id!r} inexistente.")
    if c["estado"] != "generada" or not c["n_escenarios"]:
        raise ValueError("La combinación no tiene escenarios generados — generalos primero.")
    with closing(_connect(path)) as con, con:
        con.execute("INSERT INTO config VALUES ('active_combination', ?) "
                    "ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor", (combo_id,))


# ── Adaptadores al pipeline existente ─────────────────────────────────────────
def seed_for(combo_id: str, *, fecha_inicial: str | None = None, fecha_final: str | None = None,
             tickers: list | None = None, path: Path = DB_PATH):
    """ucbatch.reader.Seed desde el seed de la combinación, con override de fechas/tickers por
    corrida (el job incremental pasa el rango del día; la reevaluación, el de la UI)."""
    from ucbatch.reader import Seed

    c = get_combination(combo_id, path=path)
    if not c:
        raise ValueError(f"Combinación {combo_id!r} inexistente.")
    s = c["seed"]

    def _f(key, default):
        try:
            return float(str(s.get(key, default)).replace(",", "."))
        except (TypeError, ValueError):
            return default
    seed = Seed(
        tickers=[t.upper() for t in (tickers or s.get("tickers") or [])],
        tipo=str(s.get("Tipo de operacion") or "CALL y PUT"),
        fecha_inicial=_norm(fecha_inicial or s.get("fecha_inicial")),
        fecha_final=_norm(fecha_final or s.get("fecha_final")),
        inversion=_f("Inversión ($)", 1000.0),
        call_pct=_f("Inversión CALL (%)", 50.0),
        put_pct=_f("Inversión PUT (%)", 50.0),
        entrada=str(s.get("Horario de entrada") or "09:30"),
        salida=str(s.get("Horario de salida") or "13:55"),
        ventana_min=_f("Ventana búsqueda contrato (max)", 0.0),
        criterio=str(s.get("Criterio de selección de contrato") or ""),
        fills=str(s.get("Modelo de fills") or ""),
        granularidad_seg=int(_f("Granularidad temporal (segundos)", 60)),
        dte=str(s.get("Vencimiento DTE") or "0 — mismo día"),
    )
    seed.display = [("Combinación", c["nombre"]), ("Tickers", ", ".join(seed.tickers)),
                    ("Rango", f"{seed.fecha_inicial} → {seed.fecha_final}")]
    return seed


def scenarios_for_batch(combo_id: str, path: Path = DB_PATH) -> list:
    """list[ucbatch.reader.Scenario] — cond con los HEADERS ORIGINALES → map_scenario los
    consume sin cambios."""
    from ucbatch.reader import Scenario

    with closing(_connect(path)) as con:
        rows = con.execute("SELECT codigo, cond_json FROM scenarios WHERE combination_id=? "
                           "ORDER BY codigo", (combo_id,)).fetchall()
    return [Scenario(id=r[0], cond=json.loads(r[1])) for r in rows]


def scenarios_df(combo_id: str, path: Path = DB_PATH):
    """DataFrame CANÓNICO de condiciones (id + nombres de loader._SCENARIO_RENAME) — drop-in de
    `bt_analysis.loader.load_template()[1]` para engine.analyze / _scenario_configs."""
    import pandas as pd
    from bt_analysis.loader import _SCENARIO_RENAME

    with closing(_connect(path)) as con:
        rows = con.execute("SELECT codigo, cond_json FROM scenarios WHERE combination_id=? "
                           "ORDER BY codigo", (combo_id,)).fetchall()
    if not rows:
        return pd.DataFrame(columns=["id"])
    recs = [{"ID": r[0], **json.loads(r[1])} for r in rows]
    df = pd.DataFrame(recs).rename(columns=_SCENARIO_RENAME)
    df = df.loc[:, ~df.columns.duplicated()]
    df["id"] = df["id"].astype(str).str.strip()
    return df


def scenario_configs(combo_id: str, path: Path = DB_PATH) -> dict:
    """{codigo: {condición canónica: valor}} — la misma forma que bt_analysis.playbook._scenario_configs
    (para condiciones del playbook persistido y el dropdown del panel), directo desde la DB."""
    df = scenarios_df(combo_id, path=path)
    if df.empty or "id" not in df.columns:
        return {}
    out: dict = {}
    for _, r in df.iterrows():
        cfg = {}
        for c in df.columns:
            if c == "id":
                continue
            v = r[c]
            if v is None or v != v or str(v).strip() == "":
                continue
            cfg[c] = str(v).strip()
        out[str(r["id"]).strip()] = cfg
    return out


# ── Migración del template legacy (import ÚNICO — después el archivo no se usa) ─
def ensure_legacy(template_path: Path = LEGACY_TEMPLATE, path: Path = DB_PATH) -> str | None:
    """Importa el template estático de 480 escenarios como la combinación «tpl480» (una sola
    vez) y la deja ACTIVA si no hay otra — así el almacén (bt_results.db, filas legacy) y el
    playbook vigente siguen siendo coherentes sin intervención. Devuelve el id (o None si el
    template no existe y no hay nada que migrar)."""
    if get_combination(LEGACY_ID, path=path):
        if active_combination(path=path) is None:
            set_active(LEGACY_ID, path=path)
        return LEGACY_ID
    if not Path(template_path).exists():
        return None
    from ucbatch import reader as _ucr

    seed_obj, scens = _ucr.read_template(str(template_path))
    seed = {"tickers": seed_obj.tickers, "Tipo de operacion": seed_obj.tipo,
            "fecha_inicial": seed_obj.fecha_inicial, "fecha_final": seed_obj.fecha_final,
            "Inversión ($)": seed_obj.inversion, "Inversión CALL (%)": seed_obj.call_pct,
            "Inversión PUT (%)": seed_obj.put_pct, "Horario de entrada": seed_obj.entrada,
            "Horario de salida": seed_obj.salida,
            "Ventana búsqueda contrato (max)": seed_obj.ventana_min,
            "Criterio de selección de contrato": seed_obj.criterio,
            "Modelo de fills": seed_obj.fills,
            "Granularidad temporal (segundos)": seed_obj.granularidad_seg,
            "Vencimiento DTE": seed_obj.dte}
    now = _now()
    with closing(_connect(path)) as con, con:
        con.execute("INSERT INTO combinations (id, nombre, archivo, creado_en, estado, "
                    "seed_json, variables_json, n_escenarios, generado_en) "
                    "VALUES (?,?,?,?,?,?,?,?,?)",
                    (LEGACY_ID, "Template 480 (legacy)", Path(template_path).name, now,
                     "generada", json.dumps(seed, ensure_ascii=False, default=str),
                     json.dumps({}, ensure_ascii=False), len(scens), now))
        con.executemany(
            "INSERT INTO scenarios VALUES (?,?,?,?)",
            [(LEGACY_ID, s.id, json.dumps({k: ("" if v is None else v) for k, v in s.cond.items()},
                                          ensure_ascii=False, default=str), now)
             for s in scens])
    if active_combination(path=path) is None:
        set_active(LEGACY_ID, path=path)
    return LEGACY_ID
