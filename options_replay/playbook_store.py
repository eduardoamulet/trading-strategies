"""Playbook persistido (Configuración → Backtesting) — la tabla día-de-la-semana OPERAR/NO OPERAR
con su escenario, métricas y CONDICIONES, evaluada sobre un rango de fechas.

Fuente de verdad de las condiciones: el TEMPLATE del batch (hoja «Backtesting scenarios»).
Pipeline de reevaluación: el MISMO del puente de la app — copiar el template con el rango/tickers
elegidos, correr `run_ucbatch.py` en una CONSOLA APARTE (480 escenarios, multiproceso) y, al
terminar, interpretar el results detallado con bt_analysis → veredicto por día a nivel CARTERA.

Persistencia: data/playbook.json (+ data/playbook_job.json mientras corre la reevaluación).
Sin Streamlit: lo consume configuracion_app.py (UI) y app.py («(playbook automático)»).

SUPUESTOS (documentados):
  · «Playbook» = veredicto por día a nivel cartera (Σganancia/Σinversión por fecha) del motor de
    interpretación, con el escenario ganador por Sharpe y sus condiciones del template.
  · «Aplicación automática por fecha» = por DÍA DE LA SEMANA del playbook guardado; el panel
    aplica UNA config por corrida, así que solo aplica cuando las filas cargadas comparten un
    único día de la semana (si no, avisa). NO OPERAR bloquea la aplicación.
  · La reevaluación tarda ~20 s por día hábil del rango (batch completo de 480 escenarios).
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).parent
PB_PATH = HERE / "data" / "playbook.json"
JOB_PATH = HERE / "data" / "playbook_job.json"
TEMPLATE_PATH = HERE.parent / "excels for backtesting" / "Backtesting_use_cases_template.xlsx"
RESULTS_DIR = HERE.parent / "resultados"

_WD_ES = {0: "Lun", 1: "Mar", 2: "Mié", 3: "Jue", 4: "Vie", 5: "Sáb", 6: "Dom"}


# ── Persistencia ─────────────────────────────────────────────────────────────
def load_playbook(path: Path = PB_PATH):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001 — sin playbook guardado todavía
        return None


def save_playbook(pb: dict, path: Path = PB_PATH) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(pb, indent=2, ensure_ascii=False), encoding="utf-8")


def scenario_for_date(pb, fecha) -> dict | None:
    """Config del playbook que corresponde a una FECHA (por su día de la semana). Devuelve la
    entrada del día ({scenario, recommendation, config, …}) o None si no hay playbook / el día
    no está cubierto (finde). El caller decide qué hacer con recommendation=NO OPERAR."""
    if not pb:
        return None
    try:
        import pandas as pd
        dia = _WD_ES.get(pd.Timestamp(str(fecha)).weekday())
    except Exception:  # noqa: BLE001
        return None
    return (pb.get("per_day") or {}).get(dia)


# ── Construcción desde un results detallado (post-batch) ─────────────────────
def build_from_results(results_path, template_path=TEMPLATE_PATH) -> dict:
    """Interpreta el results DETALLADO + template → dict del playbook persistible: rango evaluado,
    tickers, veredicto por día (escenario, WR, ROI cartera, Sharpe, n, condiciones) y el desglose
    día×ticker. Usa el mismo motor que «Interpretar resultados»."""
    import pandas as pd
    from bt_analysis import engine, loader
    from trade_plan import scenario_config_summary

    rdf, gran = loader.load_results(str(results_path))
    seed, sc = loader.load_template(str(template_path))
    rep = engine.analyze(rdf, gran, scenarios_df=sc, seed=seed)
    dow = rep.get("dow") or {}
    if not dow.get("available"):
        raise ValueError("El results no permite el análisis por día de la semana "
                         "(¿granularidad agregada?). Usá un results DETALLADO.")
    # Rango evaluado: de las FECHAS del results (no del template, que puede traer otro seed).
    _f = pd.to_datetime(rdf.get("fecha"), errors="coerce").dropna()
    from bt_analysis import playbook as _pbk
    cfgs = _pbk._scenario_configs(rep)

    per_day = {}
    for dia, i in (dow.get("per_day") or {}).items():
        cfg = cfgs.get(str(i.get("scenario") or "").strip())
        per_day[dia] = {
            "scenario": i.get("scenario"), "win_rate": i.get("win_rate"),
            "avg_roi": i.get("avg_roi"), "sharpe": i.get("sharpe"), "n": i.get("n"),
            "recommendation": i.get("recommendation"), "reason": i.get("reason", ""),
            "config": cfg, "config_txt": scenario_config_summary(cfg),
        }
    dt = rep.get("dow_ticker")
    por_ticker = dt.to_dict("records") if dt is not None and not dt.empty else []

    return {
        "evaluado_desde": str(_f.min().date()) if len(_f) else None,
        "evaluado_hasta": str(_f.max().date()) if len(_f) else None,
        "tickers": str(seed.get("Tickers") or ""),
        "generado_en": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "results_file": str(results_path),
        "n_dias": rep.get("n_days"), "n_posiciones": rep.get("n_positions"),
        "per_day": per_day, "por_ticker": por_ticker,
    }


# ── Actualización INCREMENTAL: veredicto sobre ventana rodante con decaimiento ───────────────────
HISTORY_PATH = HERE / "data" / "playbook_history.jsonl"


def _wmetrics(rois, w) -> dict:
    """Métricas PONDERADAS de una serie de ROI de cartera diarios: win-rate, media, Sharpe con
    pesos exponenciales (los días recientes votan más). `n` queda SIN ponderar (muestra honesta)."""
    sw = float(w.sum())
    if sw <= 0 or len(rois) == 0:
        return {}
    wr = float(w[rois > 0].sum() / sw * 100)
    mean = float((w * rois).sum() / sw)
    var = float((w * (rois - mean) ** 2).sum() / sw)
    std = var ** 0.5
    sharpe = (mean / std) if std > 0 else float("nan")
    return {"n": int(len(rois)), "win_rate": round(wr, 1), "avg_roi": round(mean, 3),
            "sharpe": (round(sharpe, 3) if sharpe == sharpe else None)}


def _gate(m: dict, min_n: int) -> tuple[bool, str]:
    if not m:
        return False, "sin datos"
    if m["n"] < min_n:
        return False, f"muestra insuficiente (n={m['n']} < {min_n} días)"
    # Sharpe None = desviación 0 (serie constante): si la media es positiva, es el MEJOR caso
    # (ganancia determinística), no un rechazo.
    sh_ok = (m["sharpe"] > 0) if m["sharpe"] is not None else (m["avg_roi"] > 0)
    ok = m["win_rate"] > 55 and m["avg_roi"] > 0 and sh_ok
    return ok, ("ventaja ponderada: WR>55% · ROI cartera>0 · Sharpe>0" if ok
                else "sin ventaja (WR≤55% o ROI≤0 o Sharpe≤0)")


def _rank(m: dict) -> float:
    """Clave de orden para elegir el mejor escenario: Sharpe ponderado; con desviación 0
    (Sharpe None), una media positiva rankea arriba y una negativa abajo."""
    if m.get("sharpe") is not None:
        return m["sharpe"]
    return 9e9 if (m.get("avg_roi") or 0) > 0 else -9e9


def compute_weighted_verdict(df, window_dates: list[str], *, half_life: int = 35,
                             min_n: int = 10) -> tuple[dict, list]:
    """PURO: filas canónicas (formato loader) + fechas de la ventana → (per_day, por_ticker) con
    pesos w = 0.5^(días_hábiles_atrás / half_life). ROI del día = Σpnl/Σinv (nivel cartera);
    por_ticker usa el Σ del ticker ese día. El mejor escenario del día se elige por Sharpe
    ponderado; el gate exige n (sin ponderar) ≥ min_n."""
    import numpy as np
    from bt_analysis import detailed as det

    d = det.prepare(df[df["fecha"].astype(str).str[:10].isin(window_dates)].copy())
    if d.empty:
        return {}, []
    # días hábiles atrás DENTRO de la ventana (0 = el más reciente) → peso exponencial
    ago = {f: i for i, f in enumerate(reversed(window_dates))}
    d["_fstr"] = d["date"].dt.strftime("%Y-%m-%d")
    d["_w"] = d["_fstr"].map(ago).map(lambda a: 0.5 ** (a / half_life) if a == a else np.nan)
    d = d.dropna(subset=["_w"])

    def _daily(gr, keys):
        g = gr.groupby(keys + ["_fstr"]).agg(pnl=("pnl", "sum"), inv=("inv", "sum"),
                                             w=("_w", "first"))
        g["roi"] = g["pnl"] / g["inv"].replace(0, np.nan) * 100.0
        return g.dropna(subset=["roi"])

    per_day: dict = {}
    g_cart = _daily(d, ["weekday", "id"])
    for wd in det._WD_ORDER:
        try:
            sub = g_cart.xs(wd, level="weekday")
        except KeyError:
            continue
        cand = []
        for sid, gg in sub.groupby(level="id"):
            m = _wmetrics(gg["roi"], gg["w"])
            if m:
                cand.append({"scenario": sid, **m})
        if not cand:
            per_day[wd] = {"recommendation": "NO OPERAR", "reason": "sin datos ese día"}
            continue
        best = max(cand, key=_rank)
        ok, reason = _gate(best, min_n)
        per_day[wd] = {**best, "recommendation": "OPERAR" if ok else "NO OPERAR",
                       "reason": reason}

    por_ticker: list = []
    g_tk = _daily(d, ["ticker", "weekday", "id"])
    for (tk, wd), sub in g_tk.groupby(level=["ticker", "weekday"]):
        cand = []
        for sid, gg in sub.groupby(level="id"):
            m = _wmetrics(gg["roi"], gg["w"])
            if m:
                cand.append({"scenario": sid, **m})
        if not cand:
            continue
        best = max(cand, key=_rank)
        ok, _ = _gate(best, min_n)
        por_ticker.append({"Ticker": tk, "Día": wd, "Escenario": best["scenario"],
                           "Win Rate %": best["win_rate"], "ROI %": best["avg_roi"],
                           "Sharpe": best["sharpe"], "n": best["n"],
                           "Recomendación": "OPERAR" if ok else "NO OPERAR"})
    por_ticker.sort(key=lambda r: (r["Ticker"], det._WD_ORDER.index(r["Día"])))
    return per_day, por_ticker


def build_from_store(*, window_days: int = 60, half_life: int = 35, min_n: int = 10,
                     template_path=TEMPLATE_PATH) -> dict:
    """Playbook INCREMENTAL: agrega sobre el almacén (bt_store) los últimos `window_days` días
    hábiles con decaimiento exponencial → mismo dict que build_from_results (la página y el modo
    automático no cambian) + metadatos del modo incremental."""
    import bt_store
    from bt_analysis import loader as _l, playbook as _pbk
    from trade_plan import scenario_config_summary

    dates = bt_store.distinct_dates()
    if not dates:
        raise ValueError("El almacén está vacío — corré `update_playbook.py --bootstrap` o una "
                         "Reevaluación desde la página Playbook.")
    win = dates[-window_days:]
    df = bt_store.load_range(win[0], win[-1])
    per_day, por_ticker = compute_weighted_verdict(df, win, half_life=half_life, min_n=min_n)

    # Condiciones del template por escenario (misma fuente única de siempre).
    try:
        sc = _l.load_template(str(template_path))[1]
        cfgs = _pbk._scenario_configs({"joined": sc})
    except Exception:  # noqa: BLE001 — sin template: el playbook queda con IDs solamente
        cfgs = {}
    for i in per_day.values():
        cfg = cfgs.get(str(i.get("scenario") or "").strip())
        i["config"] = cfg
        i["config_txt"] = scenario_config_summary(cfg)

    cov = bt_store.coverage()
    return {
        "evaluado_desde": win[0], "evaluado_hasta": win[-1],
        "tickers": ", ".join(cov.get("tickers") or []),
        "generado_en": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "results_file": f"bt_results.db ({cov['filas']:,} filas)",
        "n_dias": len(win), "n_posiciones": int(len(df)),
        "modo": (f"incremental · ventana {len(win)} días hábiles · half-life {half_life}d · "
                 f"min_n {min_n}"),
        "ventana_dias": window_days, "half_life": half_life, "min_n": min_n,
        "cobertura": {"desde": cov["desde"], "hasta": cov["hasta"], "dias": cov["dias"]},
        "per_day": per_day, "por_ticker": por_ticker,
    }


def append_history(pb: dict, path: Path = HISTORY_PATH) -> None:
    """Snapshot COMPACTO del veredicto a un JSONL (auditoría: cómo evolucionan los veredictos
    día a día — el dato que antes se perdía en cada reevaluación)."""
    snap = {"generado_en": pb.get("generado_en"),
            "evaluado": [pb.get("evaluado_desde"), pb.get("evaluado_hasta")],
            "modo": pb.get("modo", "manual"),
            "per_day": {d: {k: i.get(k) for k in ("scenario", "recommendation", "win_rate",
                                                  "avg_roi", "sharpe", "n")}
                        for d, i in (pb.get("per_day") or {}).items()}}
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("a", encoding="utf-8") as f:
        f.write(json.dumps(snap, ensure_ascii=False) + "\n")


# ── Reevaluación en background (consola aparte, mismo patrón que el puente) ──
def _patch_template(dst: Path, fecha_ini: str, fecha_fin: str, tickers: list[str]) -> None:
    """Copia el template y escribe rango + tickers en el Data seed (fila 3)."""
    from openpyxl import load_workbook
    shutil.copy(TEMPLATE_PATH, dst)
    wb = load_workbook(dst)
    ws = wb["Data seed"]
    hdr = [c.value for c in ws[2]]
    fcols = [i + 1 for i, h in enumerate(hdr) if h and "fecha" in str(h).lower()]
    tcol = next((i + 1 for i, h in enumerate(hdr) if h and "ticker" in str(h).lower()), None)
    if len(fcols) < 2 or tcol is None:
        raise ValueError("El Data seed del template no tiene las columnas Fecha/Ticker esperadas.")
    ws.cell(row=3, column=fcols[0], value=str(fecha_ini))
    ws.cell(row=3, column=fcols[1], value=str(fecha_fin))
    ws.cell(row=3, column=tcol, value=", ".join(tickers))
    wb.save(dst)


def launch_reevaluation(fecha_ini: str, fecha_fin: str, tickers: list[str]) -> dict:
    """Prepara el template con el rango/tickers y lanza run_ucbatch.py en una CONSOLA NUEVA
    (async, multiproceso — no bloquea la app). Persiste el job en data/playbook_job.json para
    sobrevivir refresh/reinicio. Devuelve el dict del job."""
    from ucbatch import reader as _ucr, report as _ucrep, runner as _ucrun

    job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    job_dir = HERE / "data" / ".playbook_jobs" / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    tpl = job_dir / "template.xlsx"
    _patch_template(tpl, fecha_ini, fecha_fin, tickers)

    seed, scens = _ucr.read_template(str(tpl))
    if not scens:
        raise ValueError("El template no tiene escenarios.")
    out_xlsx = RESULTS_DIR / _ucrep.output_filename(seed)
    n_days = len(_ucrun.trading_days(seed.fecha_inicial, seed.fecha_final))
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    cmd = [sys.executable, str(HERE / "run_ucbatch.py"),
           "--excel", str(tpl), "--out", str(RESULTS_DIR), "--notify"]
    subprocess.Popen(cmd, creationflags=getattr(subprocess, "CREATE_NEW_CONSOLE", 0),
                     cwd=str(HERE))
    job = {"id": job_id, "started": datetime.now().strftime("%Y-%m-%d %H:%M"),
           "out": str(out_xlsx), "template": str(tpl),
           "rango": [str(fecha_ini), str(fecha_fin)], "tickers": tickers,
           "n_dias": n_days, "n_escenarios": len(scens)}
    JOB_PATH.parent.mkdir(parents=True, exist_ok=True)
    JOB_PATH.write_text(json.dumps(job, indent=2, ensure_ascii=False), encoding="utf-8")
    return job


def check_job():
    """(estado, job): estado ∈ {None (sin job), "running", "done"}. «done» = el results del job ya
    existe en resultados/ (el batch lo escribe al FINAL de la corrida)."""
    job = load_playbook(JOB_PATH)          # mismo formato json
    if not job:
        return None, None
    return ("done" if Path(job.get("out", "")).exists() else "running"), job


def finalize_job() -> dict:
    """El batch terminó: interpreta el results del job, guarda el playbook y limpia el job.
    Además INGESTA las filas al almacén incremental (bt_store) — así las reevaluaciones manuales
    también suman historia. Devuelve el playbook nuevo."""
    status, job = check_job()
    if status != "done":
        raise RuntimeError("El job de reevaluación todavía no terminó.")
    pb = build_from_results(job["out"], template_path=job.get("template", TEMPLATE_PATH))
    save_playbook(pb)
    append_history(pb)
    try:
        import bt_store
        pb["_ingestado_al_almacen"] = bt_store.ingest_results_file(job["out"])
    except Exception:  # noqa: BLE001 — la ingesta no debe romper el finalize
        pb["_ingestado_al_almacen"] = None
    JOB_PATH.unlink(missing_ok=True)
    return pb
