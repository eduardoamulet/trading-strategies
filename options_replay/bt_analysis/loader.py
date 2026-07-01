"""Carga + normalización + detección de granularidad de los files de backtesting.

Todos los reads devuelven columnas CANÓNICAS (nombres estables) para que el resto del motor no
dependa de la ortografía exacta del header (que varía: acentos, espacios). El results file se lee
por POSICIÓN (A..N en el orden fijo de `ucbatch/report.py`), robusto a headers de 1 o 3 filas.
"""
from __future__ import annotations

import pandas as pd

# Orden fijo de columnas del results file (report.py): A..N.
_RESULT_COLS = ["id", "ticker", "fecha", "inv", "ganancia", "roi_min_pct", "roi_min_usd",
                "roi_max_pct", "roi_max_usd", "roi_avg_pct", "roi_avg_usd", "n_pos", "n_neg", "n_err"]

SCENARIO = "scenario"     # 1 fila por escenario (agregado)
DETAILED = "detailed"     # 1 fila por (ticker×día×escenario)
EMPTY = "empty"           # solo IDs, sin datos


def _read_sheet(path, prefer: str) -> list:
    import openpyxl
    src = path
    if hasattr(path, "read"):
        path.seek(0)
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    name = prefer if prefer in wb.sheetnames else wb.sheetnames[0]
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    return rows


def load_results(path) -> tuple[pd.DataFrame, str]:
    """Lee el results file → (DataFrame canónico, granularidad ∈ {scenario, detailed, empty}).
    Encuentra la fila header (col A == 'ID') y toma las filas cuyo col A sea un ID (C###)."""
    rows = _read_sheet(path, "Backtesting Results")
    hdr_i = next((i for i, r in enumerate(rows)
                  if r and str(r[0]).strip().lower() == "id"), -1)
    data = []
    for r in rows[hdr_i + 1:]:
        if not r or r[0] is None:
            continue
        sid = str(r[0]).strip()
        if not (sid[:1].upper() == "C" and sid[1:].isdigit()):
            continue
        vals = list(r[:14]) + [None] * (14 - len(r[:14]))
        data.append(dict(zip(_RESULT_COLS, vals)))
    df = pd.DataFrame(data, columns=_RESULT_COLS)
    for c in _RESULT_COLS[3:]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df, _detect_granularity(df)


def _detect_granularity(df: pd.DataFrame) -> str:
    if df.empty:
        return EMPTY
    # ¿hay datos reales? (alguna Ganancia o ROI poblado)
    if df[["ganancia", "roi_avg_pct", "n_pos"]].notna().to_numpy().sum() == 0:
        return EMPTY
    tk = df["ticker"].astype(str)
    fe = df["fecha"].astype(str)
    joined = tk.str.contains(",", na=False).any() or fe.str.contains(r"→|->", regex=True, na=False).any()
    per_id = int(df.groupby("id").size().max()) if len(df) else 1
    if joined or per_id <= 1:
        return SCENARIO
    return DETAILED


# ── Template (definiciones de escenario + config global) ────────────────────────────────────────
def load_template(path) -> tuple[dict, pd.DataFrame]:
    """Devuelve (seed_global, scenarios_df). `seed_global` = config común a los 480 (tickers, horarios,
    período…). `scenarios_df` = 1 fila por ID con las CONDICIONES que varían (refuerzo, alcance,
    umbrales/stops…) ya con nombres canónicos."""
    seed_rows = _read_sheet(path, "Data seed")
    seed = {}
    if len(seed_rows) >= 3:
        hdr, val = seed_rows[1], seed_rows[2]         # r2 header, r3 valores
        seen = {}
        for h, v in zip(hdr, val):
            if not h:
                continue
            key = str(h).strip()
            key = key + "_2" if key in seen else key   # «Fecha inicial» duplicada → final
            seen[key] = True
            seed[key] = v
    sc_rows = _read_sheet(path, "Backtesting scenarios")
    # header real en la fila 2 (fila 1 = títulos de sección combinados)
    hdr_i = next((i for i, r in enumerate(sc_rows)
                  if r and str(r[0]).strip().lower() == "id"), 1)
    hdr = [str(c).strip() if c is not None else f"col{j}" for j, c in enumerate(sc_rows[hdr_i])]
    body = [r for r in sc_rows[hdr_i + 1:] if r and r[0] and str(r[0])[:1].upper() == "C"]
    sc = pd.DataFrame(body, columns=hdr[:len(body[0])] if body else hdr)
    sc = sc.rename(columns=_SCENARIO_RENAME)
    sc = sc.loc[:, ~sc.columns.duplicated()]
    if "id" in sc.columns:
        sc["id"] = sc["id"].astype(str).str.strip()
    return seed, sc


# nombres del template → canónicos (para correlación / clustering / display)
_SCENARIO_RENAME = {
    "ID": "id",
    "Aplicar refuerzo": "refuerzo",
    "Umbral pérdida refuerzo (%)": "refuerzo_umbral",
    "No. de veces a reforzar": "refuerzo_n",
    "Sin lookahead": "sin_lookahead",
    "Alcance de salida": "alcance",
    "Cerrar si Umbral ROI ticker": "ticker_roi_on",
    "Umbral ROI (%) del ticker": "ticker_roi",
    "Cerrar si Stop loss ticker": "ticker_stop_on",
    "Stop loss (%) del ticker": "ticker_stop",
    "Filtro confirmación 1ª vela": "filtro_confirmacion",
    "Cerrar si confirmación débil": "cerrar_confirmacion_debil",
    "Cuerpo mínimo anti-doji (%)": "cuerpo_min",
    "Cerrar si Umbral ROI colectivo": "col_roi_on",
    "Umbral ROI colectivo (%)": "col_roi",
    "Cerrar si Stop loss colectivo": "col_stop_on",
    "Stop loss (%) colectivo": "col_stop",
}


def load_dow(path):
    """File opcional «ID × día de la semana» (Analisis_ID_x_diasemana): 1 fila por (escenario × día)
    con n, win-rate, avg $, volatilidad, Sharpe, peor $. Devuelve DataFrame canónico o None."""
    try:
        rows = _read_sheet(path, "ID x dia")
    except Exception:
        return None
    hdr_i = next((i for i, r in enumerate(rows)
                  if r and any("día" in str(c).lower() or "dia" in str(c).lower() for c in r if c)), 0)
    hdr = [str(c).strip() if c is not None else f"c{j}" for j, c in enumerate(rows[hdr_i])]
    body = [r for r in rows[hdr_i + 1:] if r and r[0]]
    if not body:
        return None
    df = pd.DataFrame(body, columns=hdr[:len(body[0])])
    ren = {"ID rep": "id", "params": "params", "Día": "dia", "n": "n", "Win rate %": "win_rate",
           "Avg $": "avg_usd", "Volatilidad $": "vol_usd", "Sharpe": "sharpe", "Peor $": "peor_usd"}
    df = df.rename(columns={k: v for k, v in ren.items() if k in df.columns})
    for c in ("n", "win_rate", "avg_usd", "vol_usd", "sharpe", "peor_usd"):
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def join_scenarios(results: pd.DataFrame, scenarios: pd.DataFrame) -> pd.DataFrame:
    """Une los outcomes (results) con las condiciones de cada escenario (template) por ID."""
    if scenarios is None or scenarios.empty or "id" not in scenarios.columns:
        return results.copy()
    return results.merge(scenarios, on="id", how="left", suffixes=("", "_sc"))
