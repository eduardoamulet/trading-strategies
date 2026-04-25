"""Build formatted XLSX version of Listado_Compañias_Mas_Rapidas."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

SRC = Path(r"C:\Users\ROG ZEPHYRUS\OneDrive\Documents\Claude AI\Traiding\Listado_Compañias_Mas_Rapidas.csv")
OUT = Path(r"C:\Users\ROG ZEPHYRUS\OneDrive\Documents\Claude AI\Traiding\Listado_Compañias_Mas_Rapidas.xlsx")

# ---- Styles ----
ARIAL = "Arial"
BOLD   = Font(name=ARIAL, size=10, bold=True)
BODY   = Font(name=ARIAL, size=10)
TITLE  = Font(name=ARIAL, size=14, bold=True, color="FFFFFF")
SECTION = Font(name=ARIAL, size=12, bold=True, color="FFFFFF")

# Palette
TITLE_FILL   = PatternFill("solid", start_color="1F4E78")   # dark blue
HEADER_FILL  = PatternFill("solid", start_color="CCE5FF")   # light blue
# Profile fills (light shades)
A_FILL  = PatternFill("solid", start_color="D5E8D4")  # soft green
B_FILL  = PatternFill("solid", start_color="FFF2CC")  # soft yellow
C_FILL  = PatternFill("solid", start_color="FCE5CD")  # soft orange
D_FILL  = PatternFill("solid", start_color="F4CCCC")  # soft red

# Profile section headers (darker shades)
A_SECTION = PatternFill("solid", start_color="6AA84F")  # green
B_SECTION = PatternFill("solid", start_color="BF9000")  # dark yellow
C_SECTION = PatternFill("solid", start_color="E69138")  # orange
D_SECTION = PatternFill("solid", start_color="CC0000")  # red

thin = Side(style="thin", color="808080")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

PROFILE_FILL = {"A": A_FILL, "B": B_FILL, "C": C_FILL, "D": D_FILL}
PROFILE_SECTION_FILL = {"A": A_SECTION, "B": B_SECTION, "C": C_SECTION, "D": D_SECTION}


def is_section_header(row):
    """Detect PERFIL section header rows."""
    cell0 = (row[0] or "").strip()
    return cell0.startswith("PERFIL ")


def is_legend_header(row):
    return (row[0] or "").strip() == "LEYENDA DE PERFILES DE VOLATILIDAD"


def read_csv_rows(path):
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        return [row for row in reader]


rows = read_csv_rows(SRC)

wb = Workbook()
ws = wb.active
ws.title = "Compañías por Volatilidad"

# Title row
ws.merge_cells("A1:J1")
c = ws["A1"]
c.value = "LISTADO DE COMPAÑÍAS ORDENADO POR PERFIL DE VOLATILIDAD (MENOR → MAYOR)"
c.font = TITLE
c.fill = TITLE_FILL
c.alignment = CENTER
ws.row_dimensions[1].height = 28

# Column headers (row 3)
headers = ["No", "TICKER", "PERFIL", "RANGO ÓPTIMO", "MÍNIMO Y MÁXIMO",
           "FECHA DE ANÁLISIS", "BLOQUES", "SECTORS AND INDUSTRIES", "BOURSE", "COMPAÑÍA"]
for i, h in enumerate(headers):
    cell = ws.cell(row=3, column=i + 1, value=h)
    cell.font = BOLD
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = BORDER
ws.row_dimensions[3].height = 22

# Process data rows
current_row = 4
for src_row in rows[2:]:  # skip title + header in CSV
    stripped = [(v or "").strip() for v in src_row]
    # Skip fully empty rows
    if not any(stripped):
        continue

    # Detect section header
    if is_section_header(src_row):
        letter = stripped[0].replace("PERFIL ", "")[0] if stripped[0].startswith("PERFIL ") else ""
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=10)
        c = ws.cell(row=current_row, column=1, value=stripped[0])
        c.font = SECTION
        c.fill = PROFILE_SECTION_FILL.get(letter, HEADER_FILL)
        c.alignment = CENTER
        c.border = BORDER
        ws.row_dimensions[current_row].height = 22
        current_row += 1
        continue

    # Detect legend header
    if is_legend_header(src_row):
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=10)
        c = ws.cell(row=current_row, column=1, value=stripped[0])
        c.font = Font(name=ARIAL, size=11, bold=True)
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
        current_row += 1
        continue

    # Legend rows (A, B, C, D rows at the bottom)
    if len(stripped) > 0 and stripped[0] in ("A", "B", "C", "D") and len(stripped) >= 4:
        letter = stripped[0]
        ws.cell(row=current_row, column=1, value=letter).font = BOLD
        ws.cell(row=current_row, column=1).fill = PROFILE_SECTION_FILL[letter]
        ws.cell(row=current_row, column=1).font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        ws.cell(row=current_row, column=1).alignment = CENTER
        ws.cell(row=current_row, column=1).border = BORDER
        for i, val in enumerate(stripped[1:4]):
            cell = ws.cell(row=current_row, column=2 + i, value=val)
            cell.font = BODY
            cell.alignment = LEFT
            cell.border = BORDER
            cell.fill = PROFILE_FILL[letter]
        current_row += 1
        continue

    # Regular data row — at least # + TICKER + PROFILE
    if len(stripped) < 3:
        continue
    profile = stripped[2] if len(stripped) > 2 else ""
    fill = PROFILE_FILL.get(profile, None)

    for i in range(min(10, len(stripped))):
        val = stripped[i]
        # Try to convert # column to int
        if i == 0:
            try:
                val = int(val)
            except (ValueError, TypeError):
                pass
        cell = ws.cell(row=current_row, column=i + 1, value=val)
        cell.font = BODY
        cell.border = BORDER
        if i in (0, 2):   # # and PERFIL centered
            cell.alignment = CENTER
        else:
            cell.alignment = LEFT
        if fill:
            cell.fill = fill
    # Bold the TICKER column and PERFIL column
    ws.cell(row=current_row, column=2).font = BOLD
    ws.cell(row=current_row, column=3).font = Font(name=ARIAL, size=10, bold=True)
    current_row += 1

# Column widths
widths = [5, 10, 9, 18, 22, 18, 18, 40, 9, 50]
for i, w in enumerate(widths):
    ws.column_dimensions[get_column_letter(i + 1)].width = w

# Freeze header rows
ws.freeze_panes = "A4"

# Print settings
ws.sheet_view.zoomScale = 100
ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_options.horizontalCentered = True

wb.save(OUT)
print(f"Created: {OUT}")
print(f"Last data row: {current_row - 1}")
